"""Assemble the figure-first deck from slides.json; runs in the isolated presentation environment.

Usage: python -I deck.py <figures stage directory>
Writes deck.pptx, deck.pdf (rendered from the PPTX by LibreOffice, so the two match) and one PNG per PDF
page under qa_pages/ for visual inspection.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
TITLE_HEIGHT = Inches(0.75)


def text_box(slide, left, top, width, height, text, size, bold=False):
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    return frame


def build(stage: Path) -> None:
    spec = json.loads((stage / "slides.json").read_text())
    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDTH, HEIGHT
    blank = deck.slide_layouts[6]
    slide = deck.slides.add_slide(blank)
    text_box(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.4), spec["title"], 36, bold=True)
    text_box(slide, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.8), spec["subtitle"], 20)
    for item in spec["slides"]:
        slide = deck.slides.add_slide(blank)
        text_box(slide, Inches(0.4), Inches(0.1), Inches(12.5), TITLE_HEIGHT, item["title"], 28, bold=True)
        image = stage / item["image"]
        available_height = HEIGHT - TITLE_HEIGHT - Inches(0.1)
        picture = slide.shapes.add_picture(str(image), 0, TITLE_HEIGHT, height=available_height)
        if picture.width > WIDTH:
            scale = WIDTH / picture.width
            picture.width, picture.height = WIDTH, Emu(int(picture.height * scale))
        picture.left = int((WIDTH - picture.width) / 2)
        notes = slide.notes_slide.notes_text_frame
        notes.text = "\n\n".join([item["caption"], item["notes"], "Plotted values: " + "; ".join(item["source_data"])])
    deck.save(stage / "deck.pptx")
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(stage), str(stage / "deck.pptx")], check=True, capture_output=True, timeout=600)
    pages = stage / "qa_pages"
    if pages.exists():
        shutil.rmtree(pages)
    pages.mkdir()
    with pymupdf.open(stage / "deck.pdf") as pdf:
        if pdf.page_count != len(spec["slides"]) + 1:
            raise SystemExit(f"PDF has {pdf.page_count} pages for {len(spec['slides']) + 1} slides")
        for number, page in enumerate(pdf, start=1):
            page.get_pixmap(dpi=110).save(pages / f"page_{number:02d}.png")


if __name__ == "__main__":
    build(Path(sys.argv[1]))
